# -*- coding: utf-8 -*-
"""学习文档库模块路由（卡07）。

上传 docx/pdf/xlsx/xls/md/txt/png/jpg/jpeg（单文件 ≤50MB）→ 存 data/docs/{uuid}{后缀}
→ 按类型抽取文本摘要写 documents.summary（md/txt 全文截 2000 字、pdf 前 10 页、
docx 段落拼接、xlsx 列 sheet 名+每 sheet 前 50 行、图片留空）。

分类：categories scope='doc'，支持多级 parent_id；删除分类级联删除子孙分类，
其下文档 category_id 置空（归未分类）——对应 PRD 9.1 合规风险"上传者可删除"。

埋点：doc_upload（doc_type）、doc_search（keyword, result_cnt）写 event_logs。
"""
import json
import uuid
from pathlib import Path

from fastapi import APIRouter, File, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel

from app.config import BASE_DIR, DOCS_DIR
from app.db import get_conn

router = APIRouter(prefix="/api/docs", tags=["docs"])

MAX_DOC_BYTES = 50 * 1024 * 1024  # 单文件 50MB 上限
ALLOWED_DOC_SUFFIX = {".docx", ".pdf", ".xlsx", ".xls", ".md", ".txt", ".png", ".jpg", ".jpeg"}

# md/txt 摘要截断长度；其他文本类兜底上限（防损坏文件产生超长摘要）
SUMMARY_LIMIT = 2000
SUMMARY_HARD_LIMIT = 20000

# /raw 内联展示用的 Content-Type（不设 disposition，供 iframe/img 直接引用）
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
}


def _ok(data=None):
    return {"code": 0, "data": data}


def _fail(msg: str, code: int = 1):
    """业务失败：HTTP 200 + code!=0，避免前端拿到 500。"""
    return JSONResponse(status_code=200, content={"code": code, "msg": msg})


def _log_event(conn, event: str, params: dict) -> None:
    conn.execute(
        "INSERT INTO event_logs (event, params_json) VALUES (?, ?)",
        (event, json.dumps(params, ensure_ascii=False)),
    )


# ==================== 文本摘要抽取 ====================

def _read_text_file(path: Path) -> str:
    """md/txt：直接读全文（UTF-8 失败退 GBK），截断 2000 字。"""
    raw = path.read_bytes()
    for enc in ("utf-8", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            text = raw.decode("utf-8", errors="replace")
    return text.strip()[:SUMMARY_LIMIT]


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages[:10]:  # 前 10 页文本
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts).strip()[:SUMMARY_HARD_LIMIT]


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text and p.text.strip())
    return text.strip()[:SUMMARY_HARD_LIMIT]


def _extract_xlsx(path: Path) -> str:
    """列出 sheet 名 + 每个 sheet 前 50 行文本（单元格用 | 拼接）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"【{ws.title}】")
        for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
        parts.append("")
    wb.close()
    return "\n".join(parts).strip()[:SUMMARY_HARD_LIMIT]


def _extract_summary(path: Path, suffix: str) -> str:
    """按类型抽取摘要；抽取失败不阻断上传（摘要留空）。"""
    try:
        if suffix in (".md", ".txt"):
            return _read_text_file(path)
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix in (".xlsx", ".xls"):
            return _extract_xlsx(path)
        return ""  # 图片
    except Exception:
        return ""


# ==================== 上传 ====================

@router.post("/upload")
async def upload_doc(file: UploadFile = File(...)):
    """上传文档：校验后缀与 50MB 上限 → 流式落盘 data/docs → 抽取摘要 → 写库。"""
    origin_name = file.filename or "未命名"
    suffix = Path(origin_name).suffix.lower()
    if suffix not in ALLOWED_DOC_SUFFIX:
        return _fail(
            f"不支持的文件类型「{suffix or '无后缀'}」，仅支持 .docx / .pdf / .xlsx / .xls"
            f" / .md / .txt / .png / .jpg / .jpeg"
        )

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    path = DOCS_DIR / f"{uuid.uuid4().hex}{suffix}"
    size = 0
    try:
        with open(path, "wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_DOC_BYTES:
                    raise ValueError("文件超过 50MB 上限，请压缩或拆分后再上传")
                out.write(chunk)
    except ValueError as e:
        path.unlink(missing_ok=True)
        return _fail(str(e))
    except OSError:
        path.unlink(missing_ok=True)
        return _fail("文件保存失败，请检查 data/docs 目录写入权限")
    finally:
        await file.close()

    summary = _extract_summary(path, suffix)
    doc_type = suffix.lstrip(".")
    relative = path.relative_to(BASE_DIR).as_posix()
    try:
        with get_conn() as conn:
            cur = conn.execute(
                "INSERT INTO documents (name, path, doc_type, summary) VALUES (?, ?, ?, ?)",
                (origin_name, relative, doc_type, summary),
            )
            doc_id = cur.lastrowid
            _log_event(conn, "doc_upload",
                       {"doc_id": doc_id, "doc_type": doc_type, "size_bytes": size})
    except Exception:
        path.unlink(missing_ok=True)
        return _fail("文档信息写入数据库失败，已回滚")

    return _ok({
        "id": doc_id,
        "name": origin_name,
        "doc_type": doc_type,
        "size_bytes": size,
        "summary_len": len(summary),
    })


# ==================== 分类目录（scope='doc'，多级） ====================

def _get_doc_category(conn, cat_id: int):
    return conn.execute(
        "SELECT id, name, parent_id FROM categories WHERE id = ? AND scope = 'doc'",
        (cat_id,),
    ).fetchone()


@router.get("/categories")
def doc_categories():
    """文档分类树（多级）+ 各分类文档数 + 未分类数。"""
    with get_conn() as conn:
        cats = conn.execute(
            "SELECT id, name, parent_id FROM categories WHERE scope='doc' ORDER BY id"
        ).fetchall()
        counts = {
            r["category_id"]: r["c"]
            for r in conn.execute(
                """SELECT category_id, COUNT(*) AS c FROM documents
                   WHERE category_id IS NOT NULL GROUP BY category_id"""
            )
        }
        uncategorized = conn.execute(
            "SELECT COUNT(*) AS c FROM documents WHERE category_id IS NULL"
        ).fetchone()["c"]

    nodes = {c["id"]: {"id": c["id"], "name": c["name"], "parent_id": c["parent_id"],
                       "count": counts.get(c["id"], 0), "children": []}
             for c in cats}
    tree = []
    for c in cats:
        node = nodes[c["id"]]
        if c["parent_id"] is not None and c["parent_id"] in nodes:
            nodes[c["parent_id"]]["children"].append(node)
        else:
            tree.append(node)
    return _ok({"categories": tree, "uncategorized": uncategorized,
                "total": sum(counts.values()) + uncategorized})


class CatCreateBody(BaseModel):
    name: str
    parent_id: int | None = None


@router.post("/categories")
def create_doc_category(body: CatCreateBody):
    """新建文档分类（parent_id 为空建根分类，多级 parent_id）。"""
    name = body.name.strip()
    if not name:
        return _fail("分类名称不能为空")
    with get_conn() as conn:
        if body.parent_id is not None:
            parent = _get_doc_category(conn, body.parent_id)
            if not parent:
                return _fail(f"父分类不存在（id={body.parent_id}）")
        cur = conn.execute(
            "INSERT INTO categories (name, scope, parent_id) VALUES (?, 'doc', ?)",
            (name, body.parent_id),
        )
        cat_id = cur.lastrowid
    return _ok({"id": cat_id, "name": name, "scope": "doc", "parent_id": body.parent_id})


class CatRenameBody(BaseModel):
    name: str


@router.put("/categories/{cat_id}")
def rename_doc_category(cat_id: int, body: CatRenameBody):
    """重命名文档分类。"""
    name = body.name.strip()
    if not name:
        return _fail("分类名称不能为空")
    with get_conn() as conn:
        if not _get_doc_category(conn, cat_id):
            return _fail(f"分类不存在（id={cat_id}）")
        conn.execute("UPDATE categories SET name = ? WHERE id = ?", (name, cat_id))
    return _ok({"id": cat_id, "name": name})


@router.delete("/categories/{cat_id}")
def delete_doc_category(cat_id: int):
    """删除分类：级联删除子孙分类，其下文档 category_id 置空（归未分类）。"""
    with get_conn() as conn:
        if not _get_doc_category(conn, cat_id):
            return _fail(f"分类不存在（id={cat_id}）")
        # BFS 收集子孙分类 id（多级）
        ids = [cat_id]
        frontier = [cat_id]
        while frontier:
            marks = ",".join("?" for _ in frontier)
            children = [r["id"] for r in conn.execute(
                f"SELECT id FROM categories WHERE parent_id IN ({marks})", frontier
            ).fetchall()]
            children = [c for c in children if c not in ids]
            ids.extend(children)
            frontier = children
        marks = ",".join("?" for _ in ids)
        reset = conn.execute(
            f"UPDATE documents SET category_id = NULL WHERE category_id IN ({marks})", ids
        ).rowcount
        conn.execute(f"DELETE FROM categories WHERE id IN ({marks})", ids)
    return _ok({"id": cat_id, "removed_categories": len(ids), "docs_reset": reset})


# ==================== 文档列表 / 检索 ====================

@router.get("")
def list_docs(page: int = 1, page_size: int = 20, category_id: int = None):
    """分页文档列表：category_id 筛选（-1 表示未分类），带分类名。"""
    if page < 1 or page_size < 1 or page_size > 100:
        return _fail("分页参数不合法（page>=1，1<=page_size<=100）")
    where, params = [], []
    if category_id is not None:
        if category_id == -1:
            where.append("d.category_id IS NULL")
        else:
            where.append("d.category_id = ?")
            params.append(category_id)
    where_sql = ("WHERE " + " AND ".join(where)) if where else ""
    with get_conn() as conn:
        total = conn.execute(
            f"SELECT COUNT(*) AS c FROM documents d {where_sql}", params
        ).fetchone()["c"]
        rows = conn.execute(
            f"""SELECT d.id, d.name, d.doc_type, d.category_id, d.created_at,
                       c.name AS category_name
                FROM documents d
                LEFT JOIN categories c ON c.id = d.category_id
                {where_sql}
                ORDER BY d.id DESC LIMIT ? OFFSET ?""",
            params + [page_size, (page - 1) * page_size],
        ).fetchall()
    return _ok({"list": [dict(r) for r in rows], "total": total,
                "page": page, "page_size": page_size})


@router.get("/search")
def search_docs(q: str = ""):
    """关键字检索：LIKE 匹配 name + summary（正文摘要），记录 doc_search 埋点。"""
    q = (q or "").strip()
    if not q:
        return _ok({"list": [], "total": 0, "q": ""})
    like = f"%{q}%"
    with get_conn() as conn:
        rows = conn.execute(
            """SELECT d.id, d.name, d.doc_type, d.category_id, d.created_at,
                      c.name AS category_name, d.summary
               FROM documents d
               LEFT JOIN categories c ON c.id = d.category_id
               WHERE d.name LIKE ? OR d.summary LIKE ?
               ORDER BY d.id DESC""",
            (like, like),
        ).fetchall()
        items = []
        for r in rows:
            d = dict(r)
            summary = d.pop("summary") or ""
            pos = summary.find(q)
            # 命中摘要时截取上下文片段（命中词前后各 60 字）
            if pos >= 0:
                start = max(0, pos - 60)
                d["summary_snippet"] = ("…" if start > 0 else "") + \
                    summary[start:pos + len(q) + 60].replace("\n", " ")
            else:
                d["summary_snippet"] = summary[:120].replace("\n", " ")
            items.append(d)
        _log_event(conn, "doc_search", {"keyword": q, "result_cnt": len(items)})
    return _ok({"list": items, "total": len(items), "q": q})


# ==================== 文档操作 ====================

def _get_doc(conn, doc_id: int):
    return conn.execute(
        "SELECT id, name, path, doc_type FROM documents WHERE id = ?", (doc_id,)
    ).fetchone()


class DocCategoryBody(BaseModel):
    category_id: int | None = None


@router.put("/{doc_id}/category")
def move_doc_category(doc_id: int, body: DocCategoryBody):
    """移动文档到指定分类（category_id=null 移出分类归未分类）。"""
    with get_conn() as conn:
        if not _get_doc(conn, doc_id):
            return _fail(f"文档不存在（id={doc_id}）")
        cid = body.category_id
        if cid is not None:
            cat = conn.execute(
                "SELECT id, scope FROM categories WHERE id = ?", (cid,)
            ).fetchone()
            if not cat:
                return _fail(f"分类不存在（id={cid}）")
            if cat["scope"] != "doc":
                return _fail("仅可移动到文档分类")
        conn.execute(
            "UPDATE documents SET category_id = ? WHERE id = ?", (cid, doc_id)
        )
    return _ok({"id": doc_id, "category_id": cid})


@router.get("/{doc_id}/raw")
def raw_doc(doc_id: int):
    """返回原文件（内联展示，供 iframe/img 引用）。"""
    with get_conn() as conn:
        rec = _get_doc(conn, doc_id)
    if not rec:
        return _fail(f"文档不存在（id={doc_id}）")
    path = BASE_DIR / rec["path"]
    if not path.is_file():
        return _fail(f"物理文件已丢失（{rec['path']}）")
    media = MEDIA_TYPES.get("." + rec["doc_type"], "application/octet-stream")
    return FileResponse(path, media_type=media)


@router.get("/{doc_id}/preview")
def preview_doc(doc_id: int):
    """返回抽取文本（Office 类文件降级预览用）。"""
    with get_conn() as conn:
        rec = conn.execute(
            """SELECT d.id, d.name, d.doc_type, d.summary, d.created_at,
                      c.name AS category_name
               FROM documents d
               LEFT JOIN categories c ON c.id = d.category_id
               WHERE d.id = ?""",
            (doc_id,),
        ).fetchone()
    if not rec:
        return _fail(f"文档不存在（id={doc_id}）")
    return _ok(dict(rec))


@router.delete("/{doc_id}")
def delete_doc(doc_id: int):
    """删除文档（上传者可删）：删记录并同步删除物理文件。"""
    with get_conn() as conn:
        rec = _get_doc(conn, doc_id)
        if not rec:
            return _fail(f"文档不存在（id={doc_id}）")
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
    path = BASE_DIR / rec["path"]
    try:
        if path.is_file():
            path.unlink()
    except OSError:
        pass  # 物理文件缺失不阻断记录删除
    return _ok({"id": doc_id})
